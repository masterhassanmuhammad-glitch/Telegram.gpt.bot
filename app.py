# -*- coding: utf-8 -*-

import os
import io
import sys
import time
import requests
import threading

# إجبار بايثون على طباعة السجلات فوراً في Render بدون تخزين مؤقت
sys.stdout.reconfigure(line_buffering=True) if hasattr(sys.stdout, "reconfigure") else None
os.environ["PYTHONUNBUFFERED"] = "1"

from flask import Flask, request
from google import genai
from google.genai import types

# ============================================================
# Safe Import for Telegram Module
# ============================================================
import_error_msg = ""
try:
    from telegram_client import handle_telegram_command, is_telegram_command
except Exception as err_import:
    import_error_msg = str(err_import)
    print(f"ERROR IMPORTING TELEGRAM_CLIENT: {import_error_msg}", flush=True)

    def is_telegram_command(text):
        return bool(text and (text.strip().startswith("/") or text.strip().isdigit()))

    def handle_telegram_command(s_id, text, gemini_fn):
        return f"⚠️ تعذر تحميل وحدة التليجرام.\nالسبب: {import_error_msg}"


_genai_clients = {}
_genai_lock = threading.Lock()


def get_genai_client(api_key):
    with _genai_lock:
        if api_key not in _genai_clients:
            _genai_clients[api_key] = genai.Client(api_key=api_key)
        return _genai_clients[api_key]


app = Flask(__name__)

PAGE_ACCESS_TOKEN = os.getenv("PAGE_ACCESS_TOKEN")
VERIFY_TOKEN = os.getenv("VERIFY_TOKEN")

GEMINI_API_KEYS = [
    os.getenv("GEMINI_API_KEY_1"),
    os.getenv("GEMINI_API_KEY_2"),
    os.getenv("GEMINI_API_KEY_3"),
    os.getenv("GEMINI_API_KEY_4"),
    os.getenv("GEMINI_API_KEY_5")
]
GEMINI_API_KEYS = [key for key in GEMINI_API_KEYS if key]

if not GEMINI_API_KEYS:
    raise RuntimeError("No Gemini API keys configured.")

api_index = 0
api_lock = threading.Lock()


def get_next_api_key():
    global api_index
    with api_lock:
        key = GEMINI_API_KEYS[api_index]
        api_index = (api_index + 1) % len(GEMINI_API_KEYS)
        return key


MEMORY_SECONDS = 3 * 60 * 60
MAX_MEMORY_MESSAGES = 30
MAX_FILE_SIZE = 20 * 1024 * 1024

memory = {}


def clean_memory(sender_id):
    if sender_id not in memory:
        return
    now = time.time()
    valid_messages = [item for item in memory[sender_id] if now - item.get("time", 0) <= MEMORY_SECONDS]
    memory[sender_id] = valid_messages[-MAX_MEMORY_MESSAGES:]
    if not memory[sender_id]:
        del memory[sender_id]


def add_to_memory(sender_id, role, text):
    if not text:
        return
    if sender_id not in memory:
        memory[sender_id] = []
    memory[sender_id].append({"role": role, "text": text, "time": time.time()})
    clean_memory(sender_id)


def get_memory(sender_id):
    clean_memory(sender_id)
    return memory.get(sender_id, [])


NATIVE_MIME_PREFIXES = ("image/", "audio/", "video/")
NATIVE_MIME_EXACT = {"application/pdf"}

EXTENSION_MIME_MAP = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".ppt": "application/vnd.ms-powerpoint",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".mp4": "video/mp4",
}


def guess_extension_from_url(file_url):
    try:
        path = file_url.split("?")[0]
        for ext in EXTENSION_MIME_MAP:
            if path.lower().endswith(ext):
                return ext
    except Exception:
        pass
    return None


def download_file(file_url):
    try:
        response = requests.get(file_url, timeout=30)
        if response.status_code != 200:
            return None, None
        file_data = response.content
        if len(file_data) > MAX_FILE_SIZE:
            return None, None
        mime_type = response.headers.get("Content-Type", "").split(";")[0].strip().lower()
        if not mime_type or mime_type == "application/octet-stream":
            ext = guess_extension_from_url(file_url)
            mime_type = EXTENSION_MIME_MAP.get(ext, "application/octet-stream") if ext else "application/octet-stream"
        return file_data, mime_type
    except Exception as e:
        print(f"FILE DOWNLOAD ERROR: {repr(e)}", flush=True)
        return None, None


def extract_text_from_docx(file_data):
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_data))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    parts.append(row_text)
        return "\n".join(parts).strip()
    except Exception as e:
        print(f"DOCX EXTRACT ERROR: {repr(e)}", flush=True)
        return None


def extract_text_from_pptx(file_data):
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_data))
        parts = []
        for i, slide in enumerate(prs.slides, start=1):
            slide_lines = [f"-- Slide {i} --"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            slide_lines.append(line)
            parts.append("\n".join(slide_lines))
        return "\n\n".join(parts).strip()
    except Exception as e:
        print(f"PPTX EXTRACT ERROR: {repr(e)}", flush=True)
        return None


def extract_text_from_xlsx(file_data):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"-- Sheet: {sheet.title} --")
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join("" if cell is None else str(cell) for cell in row)
                if row_text.strip(" |"):
                    parts.append(row_text)
        return "\n".join(parts).strip()
    except Exception as e:
        print(f"XLSX EXTRACT ERROR: {repr(e)}", flush=True)
        return None


def extract_text_from_txt(file_data):
    try:
        return file_data.decode("utf-8", errors="ignore").strip()
    except Exception as e:
        print(f"TXT EXTRACT ERROR: {repr(e)}", flush=True)
        return None


TEXT_EXTRACTORS = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_text_from_docx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": extract_text_from_pptx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": extract_text_from_xlsx,
    "text/plain": extract_text_from_txt,
    "text/csv": extract_text_from_txt,
}


def process_attachment(attachment):
    attach_type = attachment.get("type")
    if attach_type not in ("image", "audio", "video", "file"):
        return None
    file_url = attachment.get("payload", {}).get("url")
    if not file_url:
        return None
    file_data, mime_type = download_file(file_url)
    if not file_data:
        return {"kind": "unsupported", "label": "تعذر تحميل الملف"}
    if mime_type in NATIVE_MIME_EXACT or mime_type.startswith(NATIVE_MIME_PREFIXES):
        return {"kind": "native", "data": file_data, "mime_type": mime_type}
    extractor = TEXT_EXTRACTORS.get(mime_type)
    if extractor:
        extracted = extractor(file_data)
        if extracted:
            return {"kind": "text", "text": extracted, "label": mime_type}
    return {"kind": "unsupported", "label": f"نوع الملف غير مدعوم ({mime_type})"}


def build_gemini_contents(sender_id, current_text=None, processed_attachments=None):
    processed_attachments = processed_attachments or []
    contents = ["""أنت مساعد ذكي داخل Facebook Messenger.
أجب باللغة التي يستخدمها المستخدم، واستخدم الذاكرة المتاحة لآخر 3 ساعات عند الحاجة.
إذا كانت هناك ملفات مرفقة قم بالتنفيذ والتحليل بدقة."""]

    for item in get_memory(sender_id):
        role = item.get("role")
        text = item.get("text")
        if text:
            prefix = "المستخدم:\n" if role == "user" else "المساعد:\n"
            contents.append(prefix + text)

    if current_text:
        contents.append(f"المستخدم الآن:\n{current_text}")

    for item in processed_attachments:
        kind = item.get("kind")
        if kind == "native":
            contents.append(
                types.Part.from_bytes(
                    data=item["data"],
                    mime_type=item["mime_type"]
                )
            )
        elif kind == "text":
            contents.append(f"محتوى نصي مستخرج من ملف ({item.get('label', '')}):\n{item.get('text', '')[:15000]}")

    return contents


def generate_gemini_response(sender_id, user_text=None, processed_attachments=None):
    api_key = get_next_api_key()
    if not api_key:
        return "عذراً، لم يتم ضبط مفاتيح Gemini API."

    try:
        client = get_genai_client(api_key)
        contents = build_gemini_contents(sender_id=sender_id, current_text=user_text, processed_attachments=processed_attachments)
        preferred_models = ["gemini-3.6-flash", "gemini-3.5-flash", "gemini-3.7-flash"]
        reply_text = None

        for model_name in preferred_models:
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents
                )
                if response and hasattr(response, "text") and response.text:
                    reply_text = response.text.strip()
                    print(f"✅ GEMINI SUCCESS using model: {model_name}", flush=True)
                    break
            except Exception as model_err:
                print(f"⚠️ Model {model_name} failed: {repr(model_err)}. Trying next...", flush=True)
                continue

        if not reply_text:
            return "عذراً، لم أتمكن من إنشاء رد."

        return reply_text
    except Exception as e:
        print(f"GEMINI ERROR: {repr(e)}", flush=True)
        return "عذراً، حدث خطأ أثناء معالجة الطلب عبر الذكاء الاصطناعي."


def send_facebook_message(recipient_id, text):
    if not text:
        return
    url = "https://graph.facebook.com/v20.0/me/messages"
    params = {"access_token": PAGE_ACCESS_TOKEN}
    
    max_length = 1000
    clean_text = text.encode("utf-8", "ignore").decode("utf-8")
    chunks = [clean_text[i:i + max_length] for i in range(0, len(clean_text), max_length)]

    for chunk in chunks:
        payload = {
            "recipient": {"id": recipient_id},
            "messaging_type": "RESPONSE",
            "message": {"text": chunk}
        }
        try:
            res = requests.post(url, params=params, json=payload, timeout=20)
            print(f"📤 FB SEND STATUS: {res.status_code}", flush=True)
            if res.status_code != 200:
                print(f"❌ FB SEND ERROR: {res.text}", flush=True)
        except Exception as e:
            print(f"❌ FB SEND EXCEPTION: {repr(e)}", flush=True)


@app.route("/", methods=["GET"])
def home():
    return "Gemini + Telegram Facebook Bot is Running!", 200


@app.route("/webhook", methods=["GET"])
def verify_webhook():
    mode = request.args.get("hub.mode")
    token = request.args.get("hub.verify_token")
    challenge = request.args.get("hub.challenge")

    if mode == "subscribe" and token == VERIFY_TOKEN:
        return challenge, 200
    return "Forbidden", 403


@app.route("/webhook", methods=["POST"])
def handle_messages():
    data = request.get_json(silent=True)
    if not data or data.get("object") != "page":
        return "Not Found", 404

    for entry in data.get("entry", []):
        for event in entry.get("messaging", []):
            message = event.get("message")
            if not message or message.get("is_echo"):
                continue

            sender_id = event.get("sender", {}).get("id")
            if not sender_id:
                continue

            user_text = message.get("text", "")
            attachments = message.get("attachments", [])

            def process_in_background(s_id, u_text, atts):
                try:
                    print(f"================================\nSENDER: {s_id}\nTEXT: {u_text}", flush=True)

                    if is_telegram_command(u_text):
                        print(f"👉 Processing Telegram Command: {u_text}", flush=True)
                        tg_reply = handle_telegram_command(s_id, u_text, generate_gemini_response)
                        if not tg_reply:
                            tg_reply = "⚠️ لم يتم إرجاع نتيجة من التليجرام."
                        send_facebook_message(s_id, tg_reply)
                        return

                    processed_atts = []
                    att_summaries = []
                    for att in atts:
                        res = process_attachment(att)
                        if res:
                            processed_atts.append(res)
                            if res["kind"] == "native":
                                att_summaries.append(res["mime_type"])
                            elif res["kind"] == "text":
                                att_summaries.append(res.get("label", "text file"))
                            elif res["kind"] == "unsupported":
                                att_summaries.append("unsupported: " + res.get("label", ""))

                    if not u_text and not processed_atts:
                        return

                    clean_memory(s_id)
                    reply_text = generate_gemini_response(sender_id=s_id, user_text=u_text, processed_attachments=processed_atts)

                    if u_text:
                        add_to_memory(s_id, "user", u_text)
                    elif processed_atts:
                        add_to_memory(s_id, "user", f"[أرسل المستخدم ملف/ملفات: {', '.join(att_summaries)}]")

                    add_to_memory(s_id, "model", reply_text)
                    send_facebook_message(s_id, reply_text)

                except Exception as ex:
                    print(f"❌ BACKGROUND PROCESS ERROR: {repr(ex)}", flush=True)
                    send_facebook_message(s_id, f"❌ حدث خطأ أثناء معالجة الطلب: {ex}")

            threading.Thread(target=process_in_background, args=(sender_id, user_text, attachments), daemon=True).start()

    return "EVENT_RECEIVED", 200


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)))
    
